"""File bytes to text blocks: list of (page or None, text).

Headings are marked with a leading "# " so the chunker can carry them.
"""
from __future__ import annotations

import io

Block = tuple[int | None, str]


def extract(name: str, data: bytes) -> list[Block]:
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return _pdf(data)
    if ext == "docx":
        return _docx(data)
    return [(None, data.decode("utf-8", errors="replace"))]


def _pdf(data: bytes) -> list[Block]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return [(n, (page.extract_text() or "").strip()) for n, page in enumerate(reader.pages, 1)]


def _docx(data: bytes) -> list[Block]:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "").lower()
        parts.append(f"# {text}" if style.startswith(("heading", "title")) else text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [(None, "\n\n".join(parts))]
