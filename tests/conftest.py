from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from groundwork.cache import MemoryCache
from groundwork.config import Settings
from groundwork.pipeline import Assistant

CORPUS = Path(__file__).resolve().parent.parent / "evals" / "corpus"


class ScriptedProvider:
    """Stands in for a model: streams a fixed reply word by word and records prompts."""
    name = "scripted"

    def __init__(self, reply: str = "Claims are due within 30 days [1]."):
        self.reply = reply
        self.prompts: list[list[dict]] = []

    async def stream(self, system, messages):
        self.prompts.append(messages)
        for word in self.reply.split(" "):
            yield word + " "

    @property
    def last_prompt(self) -> str:
        return json.dumps(self.prompts[-1]) if self.prompts else ""


def docx_bytes(sections: dict[str, str]) -> bytes:
    d = DocxDocument()
    for heading, body in sections.items():
        d.add_heading(heading, level=1)
        for para in body.split("\n\n"):
            d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def settings():
    return Settings(source="local", local_root=str(CORPUS))


@pytest.fixture
def make_assistant(settings):
    def _make(reply: str = "Claims are due within 30 days [1].", cache=None):
        provider = ScriptedProvider(reply)
        return Assistant(settings, provider, cache if cache is not None else MemoryCache()), provider
    return _make


async def collect(gen):
    return [(name, data) async for name, data in gen]
